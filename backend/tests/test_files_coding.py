"""Phase 8B: File Agent and Coding Agent - scope, verified file operations, undo, dev commands, code fixes."""

import asyncio
import json
import os
import sys
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import pytest

from conftest import FakeCoding, FakeDesktop, FakeOllama, FakeTrash, build_client, files_root
from nova.ai.ollama import parse_model_output
from nova.ai.rule_based import RuleBasedProvider
from nova.coding import editor
from nova.coding.diagnostics import parse_errors, summarize_tests
from nova.coding.projects import Project, find_project, inspect, list_projects
from nova.coding.runner import CommandSpec, Tools, plan_checks, plan_command, plan_tests, run
from nova.db import Database
from nova.files import documents as docs
from nova.files.finder import query_types, search
from nova.files.ops import FileOps
from nova.files.scope import SECRET_FILE, FileScope, ScopeError, validate_name


def arun(coro):
    return asyncio.run(coro)


def old(path: Path, content: str = "x") -> Path:
    """Write a file that looks untouched for an hour (organize skips files still being written)."""
    path.write_text(content)
    t = time.time() - 3600
    os.utime(path, (t, t))
    return path


# ------------------------------------------------------------------ scope


@pytest.fixture
def scope(tmp_path):
    home = tmp_path / "home"
    for name in ("Desktop", "Documents", "Downloads", "Pictures", "Music", "Videos"):
        (home / name).mkdir(parents=True)
    htdocs = tmp_path / "htdocs"
    (htdocs / "nova" / "data").mkdir(parents=True)
    return FileScope(lambda key: home / key.capitalize(), lambda: [str(htdocs)],
                     protected=[htdocs / "nova"], private=[htdocs / "nova" / "data"])


def test_scope_allows_only_user_folders_and_projects(scope, tmp_path):
    ok = tmp_path / "home" / "Desktop" / "a.txt"
    ok.write_text("x")
    assert scope.check(ok) == ok.resolve()
    outside = tmp_path / "elsewhere.txt"
    outside.write_text("x")
    with pytest.raises(ScopeError, match="bahar"):
        scope.check(outside)
    with pytest.raises(ScopeError, match="bahar"):  # ".." cannot escape
        scope.check(tmp_path / "home" / "Desktop" / ".." / ".." / "elsewhere.txt")


def test_scope_protects_secrets_git_nova_and_roots(scope, tmp_path):
    desktop = tmp_path / "home" / "Desktop"
    (desktop / ".env").write_text("KEY=1")
    (desktop / "id_rsa").write_text("k")
    with pytest.raises(ScopeError, match="passwords/keys"):
        scope.check(desktop / ".env")
    with pytest.raises(ScopeError, match="passwords/keys"):
        scope.check(desktop / "id_rsa")
    assert SECRET_FILE.search(".env.example") is None  # templates are fine
    git = tmp_path / "htdocs" / "shop" / ".git"
    git.mkdir(parents=True)
    (git / "config").write_text("x")
    scope.check(git / "config")  # reading is fine
    with pytest.raises(ScopeError, match=".git"):
        scope.check(git / "config", write=True)
    (tmp_path / "htdocs" / "nova" / "main.py").write_text("x")
    with pytest.raises(ScopeError, match="khud ko nahi"):
        scope.check(tmp_path / "htdocs" / "nova" / "main.py", write=True)
    with pytest.raises(ScopeError, match="data folder"):
        scope.check(tmp_path / "htdocs" / "nova" / "data")
    with pytest.raises(ScopeError, match="bunyadi"):
        scope.check_movable(desktop)


@pytest.mark.parametrize("name", ["a/b", "con", "x:y", "", "..", "trailing."])
def test_invalid_names(name):
    with pytest.raises(ScopeError):
        validate_name(name)


# ------------------------------------------------------------------ search


def test_search_ranks_exact_names_and_skips_heavy_and_private_folders(tmp_path):
    root = tmp_path / "r"
    (root / "a" / "deep").mkdir(parents=True)
    (root / "node_modules").mkdir()
    (root / "private").mkdir()
    (root / "a" / "deep" / "report.pdf").write_text("x")
    (root / "report-old.pdf").write_text("x")
    (root / "node_modules" / "report.pdf").write_text("x")
    (root / "private" / "report.pdf").write_text("x")
    (root / ".env").write_text("x")
    result = search([root], "report.pdf", exclude_dirs=frozenset({os.path.normcase(str(root / "private"))}),
                    hide=SECRET_FILE)
    assert [h.path.name for h in result.hits][:2] == ["report.pdf", "report-old.pdf"]
    assert all("node_modules" not in str(h.path) and "private" not in str(h.path) for h in result.hits)
    assert not search([root], "env", hide=SECRET_FILE).hits
    assert query_types("pdf report") == ("report", (".pdf",))


# ------------------------------------------------------------------ documents


def make_pdf(path: Path, text: str) -> None:
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for n, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % n + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(out)


def test_read_docx_xlsx_pdf_and_text(tmp_path):
    from docx import Document
    from openpyxl import Workbook

    d = Document()
    d.add_paragraph("Meeting kal 5 baje")
    d.save(str(tmp_path / "a.docx"))
    wb = Workbook()
    wb.active.append(["Naam", "Raqam"])
    wb.active.append(["Ali", 500])
    wb.save(str(tmp_path / "b.xlsx"))
    make_pdf(tmp_path / "c.pdf", "Hello NOVA PDF")
    (tmp_path / "d.py").write_text("print('x')\n", encoding="utf-8")
    assert "Meeting kal 5 baje" in docs.read_document(tmp_path / "a.docx").text
    assert "Ali\t500" in docs.read_document(tmp_path / "b.xlsx").text
    assert "Hello NOVA PDF" in docs.read_document(tmp_path / "c.pdf").text
    code = docs.read_document(tmp_path / "d.py")
    assert code.kind == "code" and code.detail == "1 lines"
    (tmp_path / "e.bin").write_bytes(b"\x00\x01")
    with pytest.raises(docs.DocumentError):
        docs.read_document(tmp_path / "e.bin")


def test_edit_keeps_crlf_and_refuses_non_utf8(tmp_path):
    f = tmp_path / "w.txt"
    f.write_bytes(b"one\r\ntwo\r\n")
    docs.append_text(f, "three")
    assert f.read_bytes() == b"one\r\ntwo\r\nthree\r\n"
    assert docs.replace_text(f, "two", "2") == 1 and b"2\r\n" in f.read_bytes()
    bad = tmp_path / "latin.txt"
    bad.write_bytes("caf\xe9".encode("cp1252"))
    with pytest.raises(docs.DocumentError, match="UTF-8"):
        docs.append_text(bad, "x")


def test_docx_append_and_replace(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("chai peeni hai")
    d.save(str(tmp_path / "n.docx"))
    docs.append_text(tmp_path / "n.docx", "aur biscuit")
    assert docs.replace_text(tmp_path / "n.docx", "chai", "coffee") == 1
    text = docs.document_text(tmp_path / "n.docx")
    assert "coffee peeni hai" in text and "aur biscuit" in text


# ------------------------------------------------------------------ file operations (direct)


@pytest.fixture
def ops(scope, tmp_path):
    db = Database(tmp_path / "ops.db")
    trash = FakeTrash(tmp_path / "_trash")
    o = FileOps(scope, db, tmp_path / "backups", lambda: tmp_path / "Reports", trash=trash, can_trash=lambda p: True)
    yield SimpleNamespace(ops=o, trash=trash, home=tmp_path / "home", tmp=tmp_path)
    db.close()


def test_create_rename_move_copy_delete_are_verified(ops):
    o, home = ops.ops, ops.home
    r = o.create_file(home / "Desktop", "notes", "pehli line")
    assert r.verification == "passed" and (home / "Desktop" / "notes.txt").read_text(encoding="utf-8") == "pehli line\n"
    with pytest.raises(ScopeError, match="pehle se"):
        o.create_file(home / "Desktop", "notes.txt")
    assert o.rename(home / "Desktop" / "notes.txt", "todo").verification == "passed"  # extension kept
    assert (home / "Desktop" / "todo.txt").exists()
    assert o.move(home / "Desktop" / "todo.txt", home / "Documents").verification == "passed"
    assert o.copy(home / "Documents" / "todo.txt").detail.endswith("todo - Copy.txt")
    assert o.copy(home / "Documents" / "todo.txt").detail.endswith("todo - Copy (2).txt")
    r = o.delete(home / "Documents" / "todo.txt")
    assert r.verification == "passed" and "Recycle Bin" in r.response and not (home / "Documents" / "todo.txt").exists()
    assert ops.trash.items == [home / "Documents" / "todo.txt"]


def test_delete_refused_without_recycle_bin(ops, scope):
    o = FileOps(scope, Database(ops.tmp / "x.db"), ops.tmp / "b", lambda: ops.tmp, trash=ops.trash,
                can_trash=lambda p: False)
    f = ops.home / "Desktop" / "a.txt"
    f.write_text("x")
    with pytest.raises(ScopeError, match="Recycle Bin nahi"):
        o.delete(f)
    assert f.exists()


def test_edit_backup_and_undo(ops):
    o = ops.ops
    f = ops.home / "Documents" / "notes.txt"
    f.write_text("chai\n", encoding="utf-8")
    assert o.edit(f, "append", text="kal meeting").verification == "passed"
    assert f.read_text(encoding="utf-8") == "chai\nkal meeting\n"
    assert o.edit(f, "replace", old="chai", new="coffee").verification == "passed"
    assert o.undo().verification == "passed" and f.read_text(encoding="utf-8") == "chai\nkal meeting\n"
    assert o.undo().verification == "passed" and f.read_text(encoding="utf-8") == "chai\n"
    assert "baqi nahi" in o.undo().response


def test_undo_refuses_to_overwrite_later_changes(ops):
    o = ops.ops
    f = ops.home / "Documents" / "notes.txt"
    f.write_text("a\n", encoding="utf-8")
    o.edit(f, "append", text="b")
    f.write_text("user changed it\n", encoding="utf-8")
    r = o.undo()
    assert r.executed is False and "mit jayega" in r.response
    assert f.read_text(encoding="utf-8") == "user changed it\n"


def test_organize_plan_skips_shortcuts_downloads_and_recent_and_undo_reverses(ops):
    o, dl = ops.ops, ops.home / "Downloads"
    for name in ("a.jpg", "b.pdf", "c.zip", "d.crdownload", "e.lnk", "f.unknownext"):
        old(dl / name)
    (dl / "fresh.png").write_text("x")  # changed just now: may still be downloading
    plan = o.plan_organize(dl)
    moved = {s.name: d.parent.name for s, d in plan.moves}
    assert moved == {"a.jpg": "Images", "b.pdf": "Documents", "c.zip": "Archives", "f.unknownext": "Other"}
    assert plan.skipped == 3
    assert "Images/" in o.organize_preview(plan)
    assert o.organize(plan).verification == "passed"
    assert (dl / "Images" / "a.jpg").exists() and (dl / "e.lnk").exists()
    assert o.undo().verification == "passed"
    assert (dl / "a.jpg").exists() and not (dl / "Images").exists()


def test_projects_are_never_organized(ops, scope, tmp_path):
    with pytest.raises(ScopeError, match="code toot"):
        ops.ops.plan_organize(tmp_path / "htdocs")


def test_folder_report_is_saved(ops):
    old(ops.home / "Pictures" / "a.jpg", "x" * 2000)
    r = ops.ops.report(ops.home / "Pictures")
    assert r.verification == "passed" and Path(r.detail).read_text(encoding="utf-8").count("Images") >= 1


# ------------------------------------------------------------------ through the API (rules brain)


def cmd(client, text):
    return client.post("/api/command", json={"text": text}).json()


def ask(client, text, timeout=10.0):
    results = []
    t = threading.Thread(target=lambda: results.append(cmd(client, text)))
    t.start()
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        pending = client.get("/api/permissions/pending").json()
        if pending:
            return t, results, pending[0]
        if not t.is_alive():
            return t, results, None
        time.sleep(0.05)
    return t, results, None


def decide(client, req, approved, remember=False):
    client.post(f"/api/permissions/{req['id']}/decision", json={"approved": approved, "remember": remember})


@pytest.fixture
def env(tmp_path):
    desktop = FakeDesktop()
    coding = FakeCoding(desktop)
    root = files_root(tmp_path)
    trash = FakeTrash(root / "_trash")
    opened: list = []
    htdocs = root / "htdocs"
    htdocs.mkdir()
    ollama = FakeOllama(models=[])
    with build_client(tmp_path, ollama, desktop, permission_timeout_s=10, trash=trash, coding=coding,
                      opened=opened) as client:
        r = client.put("/api/settings", json={"project_folders": [str(htdocs)]})
        assert r.status_code == 200, r.text
        home = root / "home"
        yield SimpleNamespace(client=client, home=home, desktop_dir=home / "Desktop", docs=home / "Documents",
                              downloads=home / "Downloads", htdocs=htdocs, trash=trash, coding=coding,
                              opened=opened, tmp=tmp_path, ollama=ollama)


def last_activity(client):
    return client.get("/api/activity?limit=1").json()[0]


def test_create_file_and_folder(env):
    body = cmd(env.client, "desktop par notes.txt banao")
    assert body["executed"] and (env.desktop_dir / "notes.txt").exists() and "Verify" in body["response"]
    cmd(env.client, "desktop par todo.txt banao aur us mein chai aur biscuit likho")
    assert (env.desktop_dir / "todo.txt").read_text(encoding="utf-8") == "chai aur biscuit\n"
    cmd(env.client, "documents mein Projects naam ka folder bana do")
    assert (env.docs / "Projects").is_dir()
    assert last_activity(env.client)["verification_status"] == "passed"
    assert "pehle se" in cmd(env.client, "desktop par notes.txt banao")["response"]


def test_rename_asks_with_real_names_then_verifies(env):
    (env.desktop_dir / "notes.txt").write_text("x")
    t, results, req = ask(env.client, "notes.txt ka naam todo.txt rakh do")
    assert req and '"notes.txt" ka naam "todo.txt" rakhna' in req["question"] and req["max_risk"] == "medium"
    decide(env.client, req, True)
    t.join()
    assert (env.desktop_dir / "todo.txt").exists() and not (env.desktop_dir / "notes.txt").exists()
    assert last_activity(env.client)["permission_status"] == "approved_by_user"


def test_ambiguous_name_asks_which_one_then_ordinal_works(env):
    (env.desktop_dir / "notes.txt").write_text("desktop wali")
    (env.docs / "notes.txt").write_text("documents wali")
    body = cmd(env.client, "notes.txt parho")
    assert "kaun si" in body["response"] and "1. notes.txt" in body["response"] and not body["executed"]
    body = cmd(env.client, "doosri wali file parho")
    assert "wali" in body["response"] and body["executed"]


def test_delete_goes_to_recycle_bin_after_permission_and_is_never_remembered(env):
    (env.downloads / "report.pdf").write_text("x")
    old(env.downloads / "report-old.pdf")  # older: listed second
    listing = cmd(env.client, "report files dhoondo")["response"]
    assert "1. report.pdf" in listing and "2. report-old.pdf" in listing
    t, results, req = ask(env.client, "pehli wali ko delete karo")
    assert req["rememberable"] is False and "Recycle Bin" in req["question"]
    decide(env.client, req, True)
    t.join()
    assert not (env.downloads / "report.pdf").exists() and env.trash.items == [(env.downloads / "report.pdf").resolve()]
    assert "Restore" in results[0]["response"]


def test_big_folder_delete_is_high_risk(env):
    big = env.docs / "bigdata"
    big.mkdir()
    for n in range(120):
        (big / f"f{n}.txt").write_text("x")
    t, results, req = ask(env.client, "bigdata folder delete karo")
    assert req["max_risk"] == "high"
    decide(env.client, req, False)
    t.join()
    assert big.exists() and not env.trash.items


def test_edit_shows_preview_then_undo_restores(env):
    f = env.docs / "notes.txt"
    f.write_text("line1\n", encoding="utf-8")
    t, results, req = ask(env.client, "notes.txt mein likho: kal meeting 5 baje")
    item = req["items"][0]
    assert item["preview"] == "+ kal meeting 5 baje" and item["rememberable"] is False
    decide(env.client, req, True)
    t.join()
    assert f.read_text(encoding="utf-8") == "line1\nkal meeting 5 baje\n"
    t, results, req = ask(env.client, "pichla file kaam undo karo")
    assert "edit" in req["question"]
    decide(env.client, req, True)
    t.join()
    assert f.read_text(encoding="utf-8") == "line1\n"


def test_replace_text(env):
    f = env.docs / "menu.txt"
    f.write_text("chai\nchai\n", encoding="utf-8")
    t, results, req = ask(env.client, "menu.txt mein 'chai' ko 'coffee' se badal do")
    assert "(2 jagah)" in req["question"]
    decide(env.client, req, True)
    t.join()
    assert f.read_text(encoding="utf-8") == "coffee\ncoffee\n"


def test_organize_downloads_with_preview(env):
    for name in ("a.jpg", "b.pdf", "c.zip"):
        old(env.downloads / name)
    t, results, req = ask(env.client, "Downloads organize karo")
    assert "Images/" in req["items"][0]["preview"] and "3 files" in req["question"]
    decide(env.client, req, True)
    t.join()
    assert (env.downloads / "Images" / "a.jpg").exists() and (env.downloads / "Archives" / "c.zip").exists()


def test_context_pronoun_and_move_copy(env):
    cmd(env.client, "desktop par plan.txt banao")
    t, results, req = ask(env.client, "isko Documents mein move karo")
    assert "plan.txt" in req["question"]
    decide(env.client, req, True)
    t.join()
    assert (env.docs / "plan.txt").exists()
    body = cmd(env.client, "plan.txt ki copy banao")  # low risk: no question
    assert body["executed"] and (env.docs / "plan - Copy.txt").exists()


def test_secret_outside_and_executable_are_refused(env):
    (env.docs / "secrets.json").write_text("{}")
    body = cmd(env.client, "secrets.json parho")
    assert "passwords/keys" in body["response"] and not body["executed"]
    assert last_activity(env.client)["permission_status"] == "refused_by_nova"
    assert "bahar" in cmd(env.client, "C:\\Windows\\win.ini parho")["response"]
    (env.downloads / "setup.exe").write_bytes(b"MZ")
    body = cmd(env.client, "setup.exe kholo")
    assert "nahi chalata" in body["response"] and not env.opened


def test_open_and_read_and_list(env):
    (env.docs / "notes.txt").write_text("hello world\n", encoding="utf-8")
    body = cmd(env.client, "notes.txt kholo")
    # The fake desktop already shows "notes.txt - Notepad", so the window check passes.
    assert env.opened and env.opened[0].name == "notes.txt" and "khul gaya" in body["response"]
    assert "hello world" in cmd(env.client, "notes.txt parho")["response"]
    assert "Documents mein 1 files" in cmd(env.client, "Documents ki files dikhao")["response"]


def test_read_with_summary_uses_local_model(tmp_path):
    ollama = FakeOllama(models=["qwen3:4b"])
    ollama.reply = lambda text: {"points": ["Ye meeting ke notes hain [1]"]} if "PAGE TITLE" in text else \
        {"intents": [{"name": "unknown"}], "answer": ""}
    with build_client(tmp_path, ollama) as client:
        (files_root(tmp_path) / "home" / "Documents" / "notes.txt").write_text("Meeting at 5\n" * 30, encoding="utf-8")
        body = cmd(client, "documents ki notes.txt ka khulasa batao")
    assert "ka khulasa" in body["response"] and "Ye meeting ke notes hain" in body["response"]


def test_folder_report_and_roots_endpoint(env):
    (env.downloads / "a.pdf").write_text("x")
    body = cmd(env.client, "Downloads folder ki report banao")
    assert "Report save ho gayi" in body["response"] and body["executed"]
    roots = env.client.get("/api/files/roots").json()
    assert {r["name"] for r in roots} >= {"Desktop", "Documents", "Downloads", "htdocs"}


# ------------------------------------------------------------------ settings


@pytest.mark.parametrize("folder,error", [
    ("C:\\", "drive"), ("C:\\Windows\\System32", "Windows"), ("relative\\path", "Poora"),
    (str(Path.home() / "AppData" / "Roaming"), "AppData"), (str(Path.home()), "user folder"),
])
def test_project_folder_validation(env, folder, error):
    r = env.client.put("/api/settings", json={"project_folders": [folder]})
    assert r.status_code == 422 and error in r.text


def test_missing_project_folder_is_refused(env):
    r = env.client.put("/api/settings", json={"project_folders": [str(env.htdocs / "nope")]})
    assert r.status_code == 422 and "Folder nahi mila" in r.text


# ------------------------------------------------------------------ coding


def make_node_project(htdocs: Path, name="shop") -> Path:
    p = htdocs / name
    p.mkdir()
    (p / "package.json").write_text(json.dumps({"name": name, "scripts": {"test": "vitest run", "build": "vite build",
                                                                          "dev": "vite"},
                                                "devDependencies": {"vite": "^5", "react": "^19"}}))
    (p / "src").mkdir()
    (p / "src" / "main.jsx").write_text("console.log(1)\n")
    return p


def make_python_project(htdocs: Path, name="calc") -> Path:
    p = htdocs / name
    p.mkdir()
    (p / "app.py").write_text("print('hi'\n", encoding="utf-8")
    return p


COMPILE_ERROR = """*** Error compiling '.\\\\app.py'...
  File ".\\app.py", line 1
    print('hi'
         ^
SyntaxError: '(' was never closed
"""


def test_projects_list_inspect_and_open_in_vs_code(env):
    make_node_project(env.htdocs)
    make_python_project(env.htdocs)
    assert "2 projects mile: calc, shop" in cmd(env.client, "mere projects dikhao")["response"]
    info = cmd(env.client, "shop project ka jaiza lo")["response"]
    assert "Node.js" in info and "Vite" in info and "npm scripts: test, build, dev" in info
    body = cmd(env.client, "shop project kholo")
    assert env.coding.launched == [("Code.exe", (env.htdocs / "shop").resolve())]
    assert "VS Code mein khul gaya" in body["response"]
    assert last_activity(env.client)["verification_status"] == "passed"


def test_run_tests_asks_shows_command_and_records_test_status(env):
    make_node_project(env.htdocs)
    env.coding.runner.results["npm test"] = [(0, " Test Files  1 passed (1)\n      Tests  3 passed (3)\n")]
    t, results, req = ask(env.client, "shop project ke tests chalao")
    assert req["max_risk"] == "medium" and "npm test" in req["items"][0]["preview"]
    decide(env.client, req, True)
    t.join()
    spec = env.coding.runner.calls[0]
    assert spec.argv == ["npm.cmd", "test"] and spec.cwd == (env.htdocs / "shop").resolve()
    assert "kamyab" in results[0]["response"] and "3 passed" in results[0]["response"]
    assert last_activity(env.client)["test_status"] == "passed"


def test_check_errors_with_builtin_checks_runs_without_asking(env):
    make_python_project(env.htdocs)
    env.coding.runner.results["python -m compileall"] = [(1, COMPILE_ERROR)]
    t, results, req = ask(env.client, "calc project mein errors check karo")
    t.join()
    assert req is None  # compileall only parses the code; nothing of the project runs
    assert "app.py:1" in results[0]["response"] and "was never closed" in results[0]["response"]


def test_fix_error_shows_diff_applies_and_verifies(tmp_path):
    ollama = FakeOllama(models=["qwen3:4b"])
    ollama.reply = lambda text: {"edits": [{"find": "print('hi'", "replace": "print('hi')"}],
                                 "explanation": "print mein band bracket missing tha."} if "TASK:" in text else \
        {"intents": [{"name": "unknown"}], "answer": ""}
    desktop = FakeDesktop()
    coding = FakeCoding(desktop)
    coding.runner.results["python -m compileall"] = [(1, COMPILE_ERROR), (0, "")]
    with build_client(tmp_path, ollama, desktop, permission_timeout_s=10, coding=coding) as client:
        htdocs = files_root(tmp_path) / "htdocs"
        htdocs.mkdir()
        client.put("/api/settings", json={"project_folders": [str(htdocs)]})
        app_py = make_python_project(htdocs) / "app.py"
        cmd(client, "calc project mein errors check karo")
        t, results, req = ask(client, "error theek karo")
        item = req["items"][0]
        assert "+print('hi')" in item["preview"] and "-print('hi'" in item["preview"] and not item["rememberable"]
        decide(client, req, True)
        t.join()
        assert app_py.read_text(encoding="utf-8") == "print('hi')\n"
        assert "wo error ab nahi hai" in results[0]["response"]
        assert client.get("/api/activity?limit=1").json()[0]["verification_status"] == "passed"
        t, results, req = ask(client, "pichla file kaam undo karo")  # the backup brings the old code back
        decide(client, req, True)
        t.join()
        assert app_py.read_text(encoding="utf-8") == "print('hi'\n"


def test_code_change_that_breaks_syntax_is_rejected_before_asking(tmp_path):
    ollama = FakeOllama(models=["qwen3:4b"])
    ollama.reply = lambda text: {"edits": [{"find": "print('hi'", "replace": "print('hi'))("}],
                                 "explanation": "x"} if "TASK:" in text else {"intents": [{"name": "unknown"}], "answer": ""}
    coding = FakeCoding(FakeDesktop())
    coding.runner.results["python -m compileall"] = [(1, COMPILE_ERROR)]
    with build_client(tmp_path, ollama, coding=coding) as client:
        htdocs = files_root(tmp_path) / "htdocs"
        htdocs.mkdir()
        client.put("/api/settings", json={"project_folders": [str(htdocs)]})
        app_py = make_python_project(htdocs) / "app.py"
        cmd(client, "calc project mein errors check karo")
        t, results, req = ask(client, "error theek karo")
        t.join()
    assert req is None and "syntax kharab" in results[0]["response"]
    assert app_py.read_text(encoding="utf-8") == "print('hi'\n"


def test_file_changed_after_proposal_is_not_overwritten(tmp_path):
    ollama = FakeOllama(models=["qwen3:4b"])
    ollama.reply = lambda text: {"edits": [{"find": "x = 1", "replace": "x = 2"}], "explanation": "x badla"} \
        if "TASK:" in text else {"intents": [{"name": "unknown"}], "answer": ""}
    with build_client(tmp_path, ollama, permission_timeout_s=10) as client:
        htdocs = files_root(tmp_path) / "htdocs"
        (htdocs / "calc").mkdir(parents=True)
        client.put("/api/settings", json={"project_folders": [str(htdocs)]})
        f = htdocs / "calc" / "app.py"
        f.write_text("x = 1\n", encoding="utf-8")
        t, results, req = ask(client, "app.py mein x ko 2 kar do")  # short: understood by the rules
        assert req is not None
        f.write_text("x = 1\ny = 5\n", encoding="utf-8")  # the user edits the file meanwhile
        decide(client, req, True)
        t.join()
    assert "badal chuki" in results[0]["response"] and f.read_text(encoding="utf-8") == "x = 1\ny = 5\n"


def test_run_command_build_dev_server_install_and_refusals(env):
    make_node_project(env.htdocs)
    t, results, req = ask(env.client, "shop project mein npm run build chalao")
    assert req["max_risk"] == "medium" and "code/scripts" in req["question"]
    decide(env.client, req, True)
    t.join()
    assert env.coding.runner.calls[-1].argv == ["npm.cmd", "run", "build"]
    t, results, req = ask(env.client, "shop project ka dev server start karo")
    decide(env.client, req, True)
    t.join()
    assert env.coding.servers and "terminal window" in results[0]["response"]
    t, results, req = ask(env.client, "shop project ki dependencies install karo")
    assert req["rememberable"] is False and "Internet" in req["question"]
    decide(env.client, req, False)
    t.join()
    assert "Mojood scripts: test, build, dev" in cmd(env.client, "shop project mein npm run deploy chalao")["response"]
    assert "Git is PC par nahi mila" in cmd(env.client, "shop project mein git status chalao")["response"]


def test_explain_error_without_model_shows_the_error(env):
    make_python_project(env.htdocs)
    env.coding.runner.results["python -m compileall"] = [(1, COMPILE_ERROR)]
    cmd(env.client, "calc project mein errors check karo")
    body = cmd(env.client, "error samjhao")
    assert "Local AI available nahi" in body["response"] and "never closed" in body["response"]


def test_explanation_drops_a_copied_prompt_example(tmp_path):
    from nova.coding.agent import EXPLAIN_EXAMPLE

    ollama = FakeOllama(models=["qwen3:4b"])
    ollama.reply = lambda text: {"explanation": [f"{EXPLAIN_EXAMPLE} Variable 'total' pehle define nahi hua."],
                                 "fix": "total = 0 pehle likhein."} if "ERROR" in text else \
        {"intents": [{"name": "unknown"}], "answer": ""}
    with build_client(tmp_path, ollama) as client:
        body = cmd(client, "error samjhao: NameError: name 'total' is not defined")
    assert "Variable 'total' pehle define nahi hua." in body["response"] and EXPLAIN_EXAMPLE not in body["response"]
    assert "Hal: total = 0 pehle likhein." in body["response"]


def test_unknown_project_and_no_project(env):
    assert "naam ka project nahi mila" in cmd(env.client, "xyz project kholo")["response"]
    assert "Kaun sa project" in cmd(env.client, "tests chalao")["response"]


# ------------------------------------------------------------------ coding units


def test_find_project_tiers(tmp_path):
    projects = [Project(n, tmp_path / n) for n in ("kn softic", "kn softic main app", "shop_app", "todo app", "todo-app")]
    assert find_project("kn softic", projects).name == "kn softic"
    assert find_project("shop app", projects).name == "shop_app"
    assert isinstance(find_project("todo app", projects), list)  # "todo app" and "todo-app" look the same
    assert find_project("shp app", projects).name == "shop_app"  # close spelling
    assert find_project("nothing", projects) is None


def test_list_projects_treats_marked_root_as_one_project(tmp_path):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    assert [p.name for p in list_projects([tmp_path])] == ["a", "b"]
    (tmp_path / "package.json").write_text("{}")
    assert [p.path for p in list_projects([tmp_path])] == [tmp_path]


def test_command_planning_is_allowlisted(tmp_path):
    p = make_node_project(tmp_path)
    info = inspect(p)
    tools = Tools(npm="npm.cmd", python="python.exe", git="git.exe")
    assert plan_tests(info, tools).argv == ["npm.cmd", "test"]
    assert plan_command(info, "build", tools).argv == ["npm.cmd", "run", "build"]
    assert plan_command(info, "dev server", tools).long_running is True
    assert plan_command(info, "git status", tools).argv == ["git.exe", "status", "--short", "--branch"]
    assert "sirf parhne wali" in plan_command(info, "git push", tools)
    assert "script nahi" in plan_command(info, "rm -rf /", tools)
    assert plan_command(info, "install", tools).network is True
    py = make_python_project(tmp_path)
    (py / "requirements.txt").write_text("requests\n")
    assert "virtual environment" in plan_command(inspect(py), "install", tools)
    assert plan_checks(inspect(py), tools)[0].runs_project_code is False


def test_runner_runs_real_commands_with_exit_code_and_timeout(tmp_path):
    spec = CommandSpec([sys.executable, "-c", "print('hello nova'); raise SystemExit(3)"], tmp_path, "py", "script", 30)
    result = run(spec)
    assert result.exit_code == 3 and "hello nova" in result.output
    slow = CommandSpec([sys.executable, "-c", "import time; time.sleep(30)"], tmp_path, "slow", "script", timeout=1)
    started = time.monotonic()
    result = run(slow)
    assert result.timed_out and time.monotonic() - started < 15
    (tmp_path / "good.py").write_text("x = 1\n")
    (tmp_path / "bad.py").write_text("x = (\n")
    per_file = CommandSpec([sys.executable, "-m", "py_compile"], tmp_path, "py_compile", "check", 60,
                           files=[tmp_path / "good.py", tmp_path / "bad.py"])
    result = run(per_file)
    assert result.exit_code == 1 and result.failed_files == [str(tmp_path / "bad.py")]


def test_parse_errors_from_common_tools(tmp_path):
    out = "\n".join([
        "FAILED tests/test_calc.py::test_add - assert 1 == 2",
        "src/app.ts(12,5): error TS2322: Type 'string' is not assignable to type 'number'.",
        "PHP Parse error:  syntax error, unexpected '}' in C:\\site\\index.php on line 7",
        str(tmp_path / "src" / "util.js"),
        "  3:10  error  'x' is defined but never used  no-unused-vars",
    ])
    errors = parse_errors(out, tmp_path)
    where = {(Path(e.file).name, e.line, e.tool) for e in errors}
    assert ("test_calc.py", None, "pytest") in where
    assert ("app.ts", 12, "typescript") in where
    assert ("index.php", 7, "php") in where
    assert ("util.js", 3, "eslint") in where
    assert summarize_tests("===== 3 passed, 1 failed in 0.52s =====") == "3 passed, 1 failed"
    assert summarize_tests("Tests:       1 failed, 4 passed, 5 total") == "1 failed, 4 passed, 5 total"


def test_editor_matches_whitespace_tolerantly_and_refuses_ambiguity():
    text = "def f():\n    return 1\n\ndef g():\n    return 1\n"
    assert editor.apply_edits(text, [{"find": "return 1", "replace": "return 2"}]) is None  # twice: ambiguous
    new = editor.apply_edits(text, [{"find": "def g():\nreturn 1", "replace": "def g():\nreturn 3"}])
    assert new == "def f():\n    return 1\n\ndef g():\n    return 3\n"  # indentation restored
    assert editor.syntax_error(Path("a.py"), "x = (", Tools()) is not None
    assert editor.syntax_error(Path("a.py"), "x = 1", Tools()) is None


# ------------------------------------------------------------------ brain


@pytest.mark.parametrize("text,name,entities", [
    ("Downloads mein pdf files dhoondo", "search_files", {"query": "pdf", "location": "Downloads"}),
    ("notes.txt ka naam final.txt rakh do", "rename_file", {"target": "notes.txt", "new_name": "final.txt"}),
    ("isko Documents mein move karo", "move_file", {"target": "is", "destination": "Documents"}),
    ("pehli wali ko delete karo", "delete_file", {"target": "pehli wali"}),
    ("notes.txt mein 'chai' ko 'coffee' se badal do", "edit_file",
     {"target": "notes.txt", "edit_action": "replace", "old_text": "chai", "new_text": "coffee"}),
    ("app.py mein login function add karo", "modify_code", {"target": "app.py"}),
    ("Downloads organize karo", "organize_folder", {"location": "Downloads"}),
    ("pichla file kaam undo karo", "undo_file_op", {}),
    ("nova project kholo", "open_project", {"project": "nova"}),
    ("VS Code mein shop project kholo", "open_project", {"project": "shop"}),
    ("nova project ke tests chalao", "run_tests", {"project": "nova"}),
    ("nova project mein errors check karo", "check_errors", {"project": "nova"}),
    ("shop project mein npm run build chalao", "run_command", {"project": "shop", "command": "npm run build"}),
    ("error theek karo", "fix_error", {}),
    ("mere projects dikhao", "inspect_project", {}),
    # existing commands keep working
    ("example.com kholo", "open_website", {"url": "example.com"}),
    ("undo karo", "keyboard_shortcut", {"keys": "undo"}),
    ("Delete link par click karo", "browser_click", {"target": "Delete"}),
    ("solar energy ki report banao", "research", {"query": "solar energy"}),
    ("desktop dikhao", "window_control", {"action": "show_desktop"}),
])
def test_rules_understand_file_and_coding_commands(text, name, entities):
    intent = arun(RuleBasedProvider().understand(text)).intents[0]
    assert intent.name == name
    assert {k: intent.entities.get(k) for k in entities} == entities


@pytest.mark.parametrize("text,name,entities", [
    # Found live: a name joined to the assistant's name lost its "NOVA-" part.
    ("NOVA-Test-8B folder mein kya hai", "read_file", {"target": "NOVA-Test-8B folder"}),
    ("nova-demo-8b project kholo", "open_project", {"project": "nova-demo-8b"}),
    ("Hey NOVA, nova-demo-8b project kholo", "open_project", {"project": "nova-demo-8b"}),
])
def test_names_starting_with_nova_are_kept(text, name, entities):
    intent = arun(RuleBasedProvider().understand(text)).intents[0]
    assert intent.name == name and {k: intent.entities.get(k) for k in entities} == entities


def test_folder_names_never_fall_back_to_a_similar_project(env):
    """Found live: "Test-8B folder" fuzzy-matched the project "test web" and a file was created there."""
    (env.htdocs / "test web").mkdir()
    (env.docs / "Test-8B-old").mkdir()
    body = cmd(env.client, "Test-8B folder mein todo.txt banao")
    assert not body["executed"] and "nahi mila" in body["response"] and "Test-8B-old" in body["response"]
    assert not (env.htdocs / "test web" / "todo.txt").exists()
    (env.docs / "Test-8B").mkdir()
    body = cmd(env.client, "Test-8B folder mein todo.txt banao")
    assert body["executed"] and (env.docs / "Test-8B" / "todo.txt").exists()


def test_successful_command_shows_its_output(env):
    make_node_project(env.htdocs)
    env.coding.runner.results["npm run build"] = [(0, "vite v5\nbuilt in 1.2s\ndist/index.html 0.5 kB")]
    t, results, req = ask(env.client, "shop project mein npm run build chalao")
    decide(env.client, req, True)
    t.join()
    assert "kamyab" in results[0]["response"] and "built in 1.2s" in results[0]["response"]


def test_long_command_carrying_text_stays_with_the_rules():
    """Found live: the model handled this long sentence and dropped the text. It is long only because of the text."""
    i = arun(RuleBasedProvider().understand(
        "NOVA-Test-8B folder mein todo.txt banao aur us mein doodh aur chai likho")).intents[0]
    assert i.name == "create_file" and i.entities["text"] == "doodh aur chai" and i.confidence >= 0.8


def test_text_with_aur_is_not_split_into_two_commands():
    u = arun(RuleBasedProvider().understand("notes.txt mein likho: chai aur biscuit"))
    assert [i.name for i in u.intents] == ["edit_file"] and u.intents[0].entities["text"] == "chai aur biscuit"


def test_model_output_file_fields_are_validated():
    raw = json.dumps({"intents": [
        {"name": "rename_file", "target": "notes.txt", "new_name": "todo.txt", "location": "Desktop"},
        {"name": "edit_file", "target": "a.txt", "old_text": "  x  ", "new_text": "y"},
        {"name": "modify_code", "target": "app.py", "instruction": "rename greet to hello", "project": "calc"},
        {"name": "open_project", "project": "nova project"},
    ], "answer": ""})
    intents, _ = parse_model_output(raw, "roman_ur", "ollama:test")
    assert intents[0].entities == {"target": "notes.txt", "location": "Desktop", "new_name": "todo.txt"}
    assert intents[1].entities["edit_action"] == "replace" and intents[1].entities["old_text"] == "x"
    assert intents[2].entities["instruction"] == "rename greet to hello"
    assert intents[3].entities == {"project": "nova"}
