"""Read and edit supported documents: plain text/code, Word (.docx), PDF and Excel (read-only).

Text is edited only when it is valid UTF-8 (anything else could be corrupted by re-encoding), and the
file's newline style and BOM are kept.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

CODE_EXTS = frozenset({
    ".py", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".php", ".html", ".htm", ".css", ".scss", ".sass", ".less",
    ".vue", ".svelte", ".java", ".kt", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".rb", ".swift", ".dart",
    ".sql", ".sh", ".bat", ".cmd", ".ps1", ".lua", ".r", ".pl",
})
TEXT_EXTS = CODE_EXTS | frozenset({
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonc", ".xml", ".yml", ".yaml", ".ini", ".cfg", ".conf",
    ".toml", ".log", ".env.example", ".gitignore", ".htaccess", ".editorconfig", ".rst", ".tex", ".srt", ".properties",
})
TEXT_NAMES = frozenset({"dockerfile", "makefile", "readme", "license", ".gitignore", ".htaccess", ".editorconfig",
                        ".env.example", "procfile"})
READABLE_EXTS = TEXT_EXTS | {".docx", ".pdf", ".xlsx", ".xlsm"}
EDITABLE_EXTS = TEXT_EXTS | {".docx"}

MAX_TEXT_BYTES = 2_000_000
MAX_PDF_PAGES = 30
MAX_SHEET_ROWS = 40


class DocumentError(Exception):
    """A document NOVA cannot read or edit; the message is Roman Urdu."""


@dataclass
class DocContent:
    kind: str  # text | code | docx | pdf | xlsx
    text: str
    detail: str  # "42 lines", "3 pages", "2 sheets"
    truncated: bool


@dataclass
class TextFile:
    text: str  # with "\n" newlines
    newline: str  # "\n" or "\r\n" as found in the file
    bom: bool


def is_text(path: Path) -> bool:
    return path.suffix.lower() in TEXT_EXTS or path.name.lower() in TEXT_NAMES


def is_readable(path: Path) -> bool:
    return is_text(path) or path.suffix.lower() in READABLE_EXTS


def is_editable(path: Path) -> bool:
    return is_text(path) or path.suffix.lower() == ".docx"


def read_text_file(path: Path, *, strict: bool = True) -> TextFile:
    size = path.stat().st_size
    if size > MAX_TEXT_BYTES:
        raise DocumentError(f"\"{path.name}\" bohat bari hai ({size // 1024} KB) — NOVA 2 MB tak ki text files parhta hai.")
    raw = path.read_bytes()
    if b"\x00" in raw[:4096]:
        raise DocumentError(f"\"{path.name}\" text file nahi lagti (binary data hai).")
    bom = raw.startswith(b"\xef\xbb\xbf")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        if strict:
            raise DocumentError(f"\"{path.name}\" UTF-8 mein nahi — NOVA isay edit nahi karega (file kharab ho sakti hai).")
        text = raw.decode("cp1252", errors="replace")
    newline = "\r\n" if "\r\n" in text else "\n"
    return TextFile(text.replace("\r\n", "\n"), newline, bom)


def write_text_file(path: Path, content: TextFile) -> None:
    encoding = "utf-8-sig" if content.bom else "utf-8"
    with open(path, "w", encoding=encoding, newline=content.newline) as f:
        f.write(content.text)


def read_document(path: Path, max_chars: int = 8000) -> DocContent:
    suffix = path.suffix.lower()
    try:
        if is_text(path):
            text = read_text_file(path, strict=False).text
            lines = text.count("\n") + (1 if text and not text.endswith("\n") else 0)
            kind = "code" if suffix in CODE_EXTS else "text"
            return DocContent(kind, text[:max_chars], f"{lines} lines", len(text) > max_chars)
        if suffix == ".docx":
            return _read_docx(path, max_chars)
        if suffix == ".pdf":
            return _read_pdf(path, max_chars)
        if suffix in (".xlsx", ".xlsm"):
            return _read_xlsx(path, max_chars)
    except DocumentError:
        raise
    except Exception as exc:  # corrupt or unusual files: explain instead of crashing
        raise DocumentError(f"\"{path.name}\" parhi nahi ja saki ({type(exc).__name__}).") from exc
    raise DocumentError(f"\"{path.name}\" ki qism ({suffix or 'bina extension'}) NOVA abhi nahi parhta. "
                        "Text, code, Word (.docx), PDF aur Excel (.xlsx) parh sakta hai.")


def _read_docx(path: Path, max_chars: int) -> DocContent:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(parts)
    return DocContent("docx", text[:max_chars], f"{len(doc.paragraphs)} paragraphs", len(text) > max_chars)


def _read_pdf(path: Path, max_chars: int) -> DocContent:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise DocumentError(f"\"{path.name}\" password se band hai — NOVA isay nahi kholta.")
    pages = reader.pages
    text_parts = []
    for page in pages[:MAX_PDF_PAGES]:
        text_parts.append((page.extract_text() or "").strip())
        if sum(len(t) for t in text_parts) > max_chars:
            break
    text = "\n\n".join(t for t in text_parts if t)
    if not text:
        raise DocumentError(f"\"{path.name}\" mein parhne layak text nahi (shayad scan ki hui tasveer hai).")
    return DocContent("pdf", text[:max_chars], f"{len(pages)} pages", len(text) > max_chars or len(pages) > MAX_PDF_PAGES)


def _read_xlsx(path: Path, max_chars: int) -> DocContent:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            parts.append(f"# {ws.title}")
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_SHEET_ROWS:
                    parts.append("...")
                    break
                if any(v is not None for v in row):
                    parts.append("\t".join("" if v is None else str(v) for v in row))
    finally:
        wb.close()
    text = "\n".join(parts)
    return DocContent("xlsx", text[:max_chars], f"{len(wb.sheetnames)} sheets", len(text) > max_chars)


# ------------------------------------------------------------------ editing


def append_text(path: Path, addition: str) -> None:
    if path.suffix.lower() == ".docx":
        from docx import Document

        doc = Document(str(path))
        for line in addition.split("\n"):
            doc.add_paragraph(line)
        doc.save(str(path))
        return
    content = read_text_file(path)
    text = content.text
    if text and not text.endswith("\n"):
        text += "\n"
    content.text = text + addition.rstrip("\n") + "\n"
    write_text_file(path, content)


def replace_text(path: Path, old: str, new: str) -> int:
    """Replace every occurrence; returns how many were replaced."""
    if path.suffix.lower() == ".docx":
        return _replace_docx(path, old, new)
    content = read_text_file(path)
    count = content.text.count(old)
    if count:
        content.text = content.text.replace(old, new)
        write_text_file(path, content)
    return count


def _replace_docx(path: Path, old: str, new: str) -> int:
    from docx import Document

    doc = Document(str(path))
    count = 0
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for p in paragraphs:
        if old not in p.text:
            continue
        in_runs = 0
        for run in p.runs:  # keeps formatting when the text sits inside one run
            if old in run.text:
                in_runs += run.text.count(old)
                run.text = run.text.replace(old, new)
        if old in p.text:  # split across runs: rewrite the paragraph text (first run's formatting kept)
            total = p.text.count(old)
            text = p.text.replace(old, new)
            for run in p.runs[1:]:
                run.text = ""
            if p.runs:
                p.runs[0].text = text
            in_runs += total
        count += in_runs
    if count:
        doc.save(str(path))
    return count


def document_text(path: Path) -> str:
    """Full text used to verify an edit."""
    if path.suffix.lower() == ".docx":
        return _read_docx(path, 10_000_000).text
    return read_text_file(path, strict=False).text
